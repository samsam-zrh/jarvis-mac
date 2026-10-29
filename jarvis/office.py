"""Local office files, opened in an installed compatible application."""
import io
import re
from pathlib import Path

from jarvis import tools
from jarvis.cancellation import check_cancelled


def _destination(title, path, suffix):
    if path:
        destination = Path(path).expanduser().absolute()
        if destination.suffix.lower() != suffix:
            raise ValueError('Extension attendue : ' + suffix)
        return destination
    name = re.sub(r'[^\w À-ÿ-]', '', title).strip()[:90] or 'Document'
    return Path.home() / 'Documents' / 'Jarvis' / (name + suffix)


def _open(result, application):
    try:
        result['opening'] = tools.open_item(result['path'], application)
    except (ValueError, RuntimeError) as error:
        result['opening_error'] = str(error)
    return result


def create_and_open_document(title: str, content: str, path: str = '') -> dict:
    """Créer un DOCX avec le texte rédigé, puis l'ouvrir dans Word ou Pages. content : une ligne par paragraphe ; respecter le nombre de lignes demandé. Chemin par défaut Documents/Jarvis, aucun écrasement."""
    result = tools.create_document(str(_destination(title, path, '.docx')), title, content)
    result['content_lines'] = len(content.splitlines())
    return _open(result, 'Word')


def create_presentation(title: str, slides: list[dict[str, str]], path: str = '') -> dict:
    """Créer et ouvrir un PowerPoint PPTX dans PowerPoint ou Keynote. slides contient des objets {title: titre, content: texte avec une ligne par idée}. Chaque objet produit une diapositive. Fournir le contenu complet rédigé, 1 à 30 diapositives, titres courts et jusqu'à 6 idées courtes par diapositive. Aucun écrasement."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    if not title.strip() or not 1 <= len(slides) <= 30:
        raise ValueError('Titre et 1 à 30 diapositives requis')
    for item in slides:
        lines = item.get('content', '').splitlines()
        if not item.get('title', '').strip() or len(item['title']) > 100 or len(lines) > 6 or any(len(line) > 180 for line in lines):
            raise ValueError('Diapositive : titre ≤ 100 caractères, jusqu’à 6 lignes de 180 caractères')
    check_cancelled()
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    presentation.core_properties.title = title
    for index, item in enumerate(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(245, 247, 250)
        for text, x, y, width, height, size, color in (
            (item['title'], .8, .8, 11.7, 1.35, 34, (20, 32, 52)),
            (item.get('content', ''), .8, 2.3, 11.7, 4.2, 24, (42, 56, 74)),
            (f'{title[:70]} · {index + 1}/{len(slides)}', .8, 7.05, 11.7, .3, 10, (82, 98, 118)),
        ):
            frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height)).text_frame
            frame.word_wrap = True
            frame.text = text
            for paragraph in frame.paragraphs:
                paragraph.font.name = 'Arial'
                paragraph.font.size = Pt(size)
                paragraph.font.color.rgb = RGBColor(*color)
                paragraph.space_after = Pt(12)
    payload = io.BytesIO()
    presentation.save(payload)
    destination = original = _destination(title, path, '.pptx')
    destination.parent.mkdir(parents=True, exist_ok=True)
    for number in range(1000):
        check_cancelled()
        try:
            with destination.open('xb') as handle:
                handle.write(payload.getvalue())
            return _open({'created': True, 'path': str(destination), 'slides': len(slides)}, 'PowerPoint')
        except FileExistsError:
            destination = original.with_name(f'{original.stem} ({number + 2}).pptx')
    raise FileExistsError('Trop de fichiers portent déjà ce nom')
