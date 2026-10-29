from pathlib import Path

import pytest
from docx import Document

from jarvis.tools import create_document, read_document, validate_web_url


def test_document_accents_and_no_overwrite(tmp_path):
    path = tmp_path / "Résumé été.docx"
    first = create_document(str(path), "Étude", "Bonjour à Samuel.\n\nRésultat vérifié.", ["https://example.com/source"])
    assert first["path"] == str(path)
    assert "Samuel" in "\n".join(p.text for p in Document(path).paragraphs)
    assert "https://example.com/source" in read_document(str(path))["text"]
    second = create_document(str(path), "Autre", "Autre contenu")
    assert second["path"] != first["path"]
    assert "Samuel" in read_document(first["path"])["text"]


def test_markdown_and_invalid_extension(tmp_path):
    result = create_document(str(tmp_path / "note.md"), "Titre", "Texte français")
    assert Path(result["path"]).read_text().startswith("# Titre")
    with pytest.raises(ValueError):
        create_document(str(tmp_path / "note.exe"), "Titre", "Texte")


def test_missing_document_is_honest(tmp_path):
    with pytest.raises(FileNotFoundError):
        read_document(str(tmp_path / "absent.pdf"))


def test_empty_word_document_opens_as_valid_blank_docx(tmp_path):
    result = create_document(str(tmp_path / "Document.docx"), "Document", "")
    assert not any(p.text for p in Document(result["path"]).paragraphs)


@pytest.mark.parametrize("url", ["file:///etc/passwd", "http://127.0.0.1:11436", "http://localhost/", "http://169.254.169.254/", "javascript:alert(1)"])
def test_web_reader_rejects_non_public_targets(url):
    with pytest.raises(ValueError):
        validate_web_url(url)


def test_search_reformulates_when_first_query_has_no_results(monkeypatch):
    import ddgs
    from jarvis.tools import search_web
    calls=[]
    class Search:
        def __init__(self,**kwargs):pass
        def text(self,query,**kwargs):
            calls.append(query)
            if len(calls)==1:raise RuntimeError('No results')
            return [{'title':'Source','href':'https://example.org','body':'Real result'}]
    monkeypatch.setattr(ddgs,'DDGS',Search)
    result=search_web('comment vérifier la batterie sur Mac')
    assert calls==['comment vérifier la batterie sur Mac','verifier batterie Mac']
    assert result['results'][0]['url']=='https://example.org'


def test_open_document_exact_name_and_ambiguity(monkeypatch,tmp_path):
    from jarvis import tools
    one=tmp_path/'Budget annuel.pdf';one.write_bytes(b'document')
    other=tmp_path/'Budget provisoire.pdf';other.write_bytes(b'document')
    monkeypatch.setattr(tools,'find_files',lambda *args:{'paths':[str(one),str(other)]})
    opened=[]
    monkeypatch.setattr(tools,'open_item',lambda path:opened.append(path) or {'opened':path})
    assert tools.open_document('budget annuel',str(tmp_path))['path']==str(one)
    with pytest.raises(ValueError,match='Plusieurs documents'):
        tools.open_document('budget',str(tmp_path))
    assert opened==[str(one)]
