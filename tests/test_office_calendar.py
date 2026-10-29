import pytest
from docx import Document
from pptx import Presentation
from jarvis import office, calendar_tools


def test_fifteen_line_document_and_fallback(monkeypatch, tmp_path):
    monkeypatch.setattr(office.tools, 'open_item', lambda path, app: {'application': 'Pages', 'fallback': True})
    lines = [f'Idée numéro {i} sur l’intelligence artificielle.' for i in range(15)]
    result = office.create_and_open_document('Essai', '\n'.join(lines), str(tmp_path/'Essai.docx'))
    assert [p.text for p in Document(result['path']).paragraphs][1:] == lines
    assert result['content_lines'] == 15 and result['opening']['application'] == 'Pages'


def test_presentation_content_no_overwrite_and_failed_opener(monkeypatch, tmp_path):
    def absent(*args): raise ValueError('Application absente')
    monkeypatch.setattr(office.tools, 'open_item', absent)
    args = ('Présentation', [{'title': 'L’IA', 'content': 'Apprendre\nCréer'}, {'title': 'Limites', 'content': 'Vérifier'}], str(tmp_path/'test.pptx'))
    first = office.create_presentation(*args)
    second = office.create_presentation(*args)
    deck = Presentation(first['path'])
    assert len(deck.slides) == 2
    assert 'Apprendre\nCréer' in [s.text for s in deck.slides[0].shapes]
    assert first['path'] != second['path'] and first['created']
    assert first['opening_error'] == 'Application absente'


def test_calendar_invalid_dates_never_call_mac(monkeypatch):
    def forbidden(*args): raise AssertionError('No automation expected')
    monkeypatch.setattr(calendar_tools, '_jxa', forbidden)
    for start, end in [('2026-10-30T09:00', '2026-10-30T10:00'), ('2026-10-30T10:00+01:00', '2026-10-30T09:00+01:00'), ('invalid', 'invalid')]:
        with pytest.raises(ValueError): calendar_tools.calendar_events(start, end)
    with pytest.raises(ValueError): calendar_tools.calendar_events('2026-10-01T00:00+02:00', '2026-12-01T00:00+01:00')


def test_calendar_title_remains_data(monkeypatch):
    scripts = []
    monkeypatch.setattr(calendar_tools, '_jxa', lambda script: scripts.append(script) or {'created': True, 'id': 'fake'})
    title = 'Titre " avec retour\n et code ; throw Error(123);'
    result = calendar_tools.create_calendar_event('calendar-id', title, '2026-10-30T09:00+01:00', '2026-10-30T10:00+01:00')
    import json
    payload = scripts[0].split('const a=', 1)[1].split(', app=Application', 1)[0]
    assert json.loads(payload)['title'] == title and result['id'] == 'fake'
